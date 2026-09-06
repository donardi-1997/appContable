"""
Importador historico de inventario y ventas para La Patrona VIP.

Lee un archivo Excel con hojas mensuales (marzo-agosto 2026) y crea registros
en la tabla inventory_history. Opcionalmente actualiza stock e ideal_stock
en la tabla products con el periodo mas reciente.

Uso:
    python backend/scripts/import_historical_inventory.py <ruta_excel> [--dry-run]

Requisitos:
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

# ── Configuracion de rutas ──────────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parent
APP_DIR = BACKEND_DIR / "app"

sys.path.insert(0, str(BACKEND_DIR))

from app.models import Base, InventoryHistory, Product


# ── Constantes ──────────────────────────────────────────────────────────────
SOURCE_TAG = "historical_excel_import"

COLUMN_MAP = {
    "PRODUCTO": "product_name",
    "INVENTARIO": "opening_stock",
    "INGRESOS": "incoming_stock",
    "EXISTENCIAS": "closing_stock",
    "VENTAS": "units_sold",
    "VALOR DE VENTA": "sale_price",
    "TOTAL VENDIDO": "total_sold",
    "STOCK IDEAL": "ideal_stock",
    "PARA PEDIDO PROMEDIO": "suggested_order",
}

SKIP_KEYWORDS = {"TOTAL", "SUBTOTAL", "TOTALES", "SUMA", "GENERAL", "PROMEDIO"}


# ── Funciones de parseo ─────────────────────────────────────────────────────

def parse_number(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, Decimal):
        return value
    if not isinstance(value, str):
        return None
    s = value.strip()
    if not s:
        return None
    s = re.sub(r"[$\s]", "", s)
    s = re.sub(r"(?<=\d)[.,](?=\d{3}(\D|$))", "", s)
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def parse_sheet_date(sheet_name: str) -> datetime | None:
    name = sheet_name.strip().upper()

    month_map = {
        "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4,
        "MAYO": 5, "JUNIO": 6, "JULIO": 7, "AGOSTO": 8,
        "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
        "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6,
        "JUL": 7, "AGO": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DIC": 12,
    }

    for month_name, month_num in month_map.items():
        if name.startswith(month_name):
            rest = name[len(month_name):].strip().lstrip("-").strip()
            year = None
            if rest.isdigit() and len(rest) == 4:
                year = int(rest)
            elif rest.isdigit() and len(rest) == 2:
                year = 2000 + int(rest)
            if year is None:
                year = 2026
            try:
                return datetime(year, month_num, 1)
            except ValueError:
                continue

    match = re.match(r"^(\d{1,2})\s*[-/]\s*(\d{4})$", name)
    if match:
        month_num = int(match.group(1))
        year = int(match.group(2))
        try:
            return datetime(year, month_num, 1)
        except ValueError:
            pass

    return None


def normalize_column_name(raw: str) -> str | None:
    cleaned = raw.strip().upper()
    cleaned = (
        cleaned.replace("A", "A").replace("E", "E").replace("I", "I")
        .replace("O", "O").replace("U", "U").replace("N", "N")
    )
    return COLUMN_MAP.get(cleaned)


def normalize_product_name(name: str) -> str:
    return " ".join(name.strip().split()).upper()


def is_valid_product_row(row_data: dict) -> bool:
    name = row_data.get("product_name")
    if not name or not isinstance(name, str) or not name.strip():
        return False
    name_upper = name.strip().upper()
    for kw in SKIP_KEYWORDS:
        if name_upper.startswith(kw) or name_upper == kw:
            return False
    has_numeric = False
    for key in ("opening_stock", "incoming_stock", "closing_stock", "units_sold"):
        v = row_data.get(key)
        if v is not None:
            has_numeric = True
            break
    return has_numeric


# ── Clases de datos ─────────────────────────────────────────────────────────

@dataclass
class ImportWarning:
    sheet: str
    row: int | None
    product: str
    message: str


@dataclass
class ImportReport:
    sheets_processed: int = 0
    sheets_skipped: int = 0
    rows_read: int = 0
    rows_valid: int = 0
    rows_skipped: int = 0
    products_found: int = 0
    products_created: int = 0
    products_pending: int = 0
    histories_created: int = 0
    histories_existing: int = 0
    inventory_warnings: int = 0
    total_warnings: int = 0
    errors: int = 0
    stock_updated: int = 0
    ideal_stock_updated: int = 0
    warnings: list[dict] = field(default_factory=list)
    errors_list: list[dict] = field(default_factory=list)
    skipped_products: list[str] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    descadres_headers: list[str] = field(default_factory=list)
    descadres_rows: int = 0
    descadres_imported: int = 0

    def to_dict(self) -> dict:
        return {
            "summary": {
                "sheets_processed": self.sheets_processed,
                "sheets_skipped": self.sheets_skipped,
                "rows_read": self.rows_read,
                "rows_valid": self.rows_valid,
                "rows_skipped": self.rows_skipped,
                "products_found": self.products_found,
                "products_created": self.products_created,
                "products_pending": self.products_pending,
                "histories_created": self.histories_created,
                "histories_existing": self.histories_existing,
                "inventory_warnings": self.inventory_warnings,
                "total_warnings": self.total_warnings,
                "errors": self.errors,
                "stock_updated": self.stock_updated,
                "ideal_stock_updated": self.ideal_stock_updated,
                "descadres_rows": self.descadres_rows,
                "descadres_imported": self.descadres_imported,
            },
            "warnings": self.warnings,
            "errors": self.errors_list,
            "skipped_products": self.skipped_products,
            "skipped_sheets": self.skipped_sheets,
            "descadres_headers": self.descadres_headers,
        }


# ── Logica principal ────────────────────────────────────────────────────────

def read_excel(filepath: Path) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheets_data: dict[str, list[dict]] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        headers_raw: list[str] = []
        header_row_idx = 0
        for i, row in enumerate(rows):
            non_empty = [str(c).strip() if c is not None else "" for c in row]
            if any(h for h in non_empty):
                headers_raw = non_empty
                header_row_idx = i
                break

        if not headers_raw:
            continue

        col_mapping: dict[int, str] = {}
        for idx, raw_header in enumerate(headers_raw):
            normalized = normalize_column_name(raw_header)
            if normalized:
                col_mapping[idx] = normalized

        if "product_name" not in col_mapping.values():
            continue

        sheet_rows: list[dict] = []
        for row in rows[header_row_idx + 1:]:
            if all(c is None for c in row):
                continue
            row_data: dict[str, Any] = {}
            for col_idx, field_name in col_mapping.items():
                if col_idx < len(row):
                    val = row[col_idx]
                    if field_name == "product_name":
                        row_data[field_name] = str(val).strip() if val is not None else ""
                    else:
                        row_data[field_name] = parse_number(val)
            sheet_rows.append(row_data)

        sheets_data[sheet_name] = sheet_rows

    wb.close()
    return sheets_data


def find_or_create_product(
    db: Session,
    product_name: str,
    report: ImportReport,
) -> Product | None:
    normalized = normalize_product_name(product_name)

    all_products = db.query(Product).filter(Product.active == True).all()
    for p in all_products:
        if normalize_product_name(p.name) == normalized:
            report.products_found += 1
            return p

    for p in all_products:
        if p.name.strip().upper() == product_name.strip().upper():
            report.products_found += 1
            return p

    new_product = Product(
        name=product_name.strip(),
        stock=Decimal("0"),
        price=Decimal("0"),
        cost_price=Decimal("0"),
        minimum_stock=Decimal("0"),
        active=True,
    )
    db.add(new_product)
    db.flush()
    report.products_created += 1
    return new_product


def process_sheet(
    db: Session,
    sheet_name: str,
    rows: list[dict],
    report: ImportReport,
    dry_run: bool,
) -> list[dict]:
    period_date = parse_sheet_date(sheet_name)
    if period_date is None:
        report.sheets_skipped += 1
        report.skipped_sheets.append(sheet_name)
        report.warnings.append({
            "sheet": sheet_name,
            "row": None,
            "product": "",
            "message": f"No se pudo interpretar la fecha de la hoja '{sheet_name}'",
        })
        report.total_warnings += 1
        return []

    records: list[dict] = []

    for row_idx, row_data in enumerate(rows, start=2):
        report.rows_read += 1

        if not is_valid_product_row(row_data):
            report.rows_skipped += 1
            continue

        report.rows_valid += 1
        product_name = row_data.get("product_name", "").strip()
        if not product_name:
            report.rows_skipped += 1
            continue

        product = find_or_create_product(db, product_name, report)
        if product is None:
            report.products_pending += 1
            report.skipped_products.append(product_name)
            continue

        existing = (
            db.query(InventoryHistory)
            .filter(
                InventoryHistory.product_id == product.id,
                InventoryHistory.period_date == period_date,
                InventoryHistory.source == SOURCE_TAG,
            )
            .first()
        )
        if existing:
            report.histories_existing += 1
            continue

        opening_stock = row_data.get("opening_stock") or Decimal("0")
        incoming_stock = row_data.get("incoming_stock") or Decimal("0")
        closing_stock = row_data.get("closing_stock") or Decimal("0")
        units_sold = row_data.get("units_sold") or Decimal("0")
        sale_price = row_data.get("sale_price") or Decimal("0")
        total_sold = row_data.get("total_sold") or Decimal("0")
        ideal_stock_val = row_data.get("ideal_stock")
        suggested_order_val = row_data.get("suggested_order")

        calculated_sales = opening_stock + incoming_stock - closing_stock
        if calculated_sales != units_sold:
            diff = abs(calculated_sales - units_sold)
            report.inventory_warnings += 1
            report.total_warnings += 1
            report.warnings.append({
                "sheet": sheet_name,
                "row": row_idx,
                "product": product_name,
                "message": (
                    f"Venta calculada ({calculated_sales}) != Excel ({units_sold}). "
                    f"Diferencia: {diff}"
                ),
            })

        calculated_total = units_sold * sale_price if sale_price else Decimal("0")
        if total_sold and calculated_total and total_sold != calculated_total:
            diff = abs(calculated_total - total_sold)
            if diff > Decimal("1"):
                report.total_warnings += 1
                report.warnings.append({
                    "sheet": sheet_name,
                    "row": row_idx,
                    "product": product_name,
                    "message": (
                        f"Total vendido calculado ({calculated_total}) != Excel ({total_sold}). "
                        f"Diferencia: {diff}"
                    ),
                })

        record = {
            "product_id": product.id,
            "period_date": period_date,
            "opening_stock": opening_stock,
            "incoming_stock": incoming_stock,
            "closing_stock": closing_stock,
            "units_sold": units_sold,
            "sale_price": sale_price,
            "total_sold": total_sold or calculated_total,
            "ideal_stock": ideal_stock_val,
            "suggested_order": suggested_order_val,
            "source": SOURCE_TAG,
            "source_sheet": sheet_name,
        }
        records.append(record)

    return records


def process_descadres(
    ws: Any,
    report: ImportReport,
    dry_run: bool,
) -> None:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return

    headers: list[str] = []
    for row in rows:
        non_empty = [str(c).strip() if c is not None else "" for c in row]
        if any(h for h in non_empty):
            headers = non_empty
            break

    report.descadres_headers = [h for h in headers if h]
    data_rows = [r for r in rows if any(c is not None for c in r)]
    report.descadres_rows = max(0, len(data_rows) - 1)

    descadres_map = {
        "PRODUCTO": "product_name",
        "FECHA": "date",
        "MOTIVO": "reason",
        "OBSERVACION": "notes",
        "OBSERVACIONES": "notes",
        "DIFERENCIA": "difference",
        "CANTIDAD": "quantity",
    }

    col_mapping: dict[int, str] = {}
    for idx, h in enumerate(headers):
        h_upper = h.strip().upper()
        if h_upper in descadres_map:
            col_mapping[idx] = descadres_map[h_upper]

    if "product_name" not in col_mapping.values():
        report.warnings.append({
            "sheet": "DESCUADRES",
            "row": None,
            "product": "",
            "message": (
                f"Estructura DESCUADRES no mapeable. "
                f"Encabezados: {report.descadres_headers}. "
                f"Revision manual requerida."
            ),
        })
        report.total_warnings += 1
        return

    report.warnings.append({
        "sheet": "DESCUADRES",
        "row": None,
        "product": "",
        "message": (
            f"Hoja DESCUADRES detectada con {report.descadres_rows} filas de datos. "
            f"Encabezados: {report.descadres_headers}. "
            f"Requiere revision manual. No se importo automaticamente."
        ),
    })
    report.total_warnings += 1


def update_product_stock_and_ideal(
    db: Session,
    records: list[dict],
    report: ImportReport,
) -> None:
    if not records:
        return

    product_latest: dict[int, dict] = {}
    for rec in records:
        pid = rec["product_id"]
        pd = rec["period_date"]
        if pid not in product_latest or pd > product_latest[pid]["period_date"]:
            product_latest[pid] = rec

    for pid, rec in product_latest.items():
        product = db.query(Product).filter(Product.id == pid).first()
        if not product:
            continue

        new_stock = rec["closing_stock"]
        if product.stock != new_stock:
            product.stock = new_stock
            report.stock_updated += 1

        if rec["ideal_stock"] is not None:
            if product.ideal_stock != rec["ideal_stock"]:
                product.ideal_stock = rec["ideal_stock"]
                report.ideal_stock_updated += 1


def run_import(
    filepath: Path,
    dry_run: bool = False,
    db_engine: Any = None,
    session_factory: Any = None,
) -> ImportReport:
    from app.db import SessionLocal as _SessionLocal, engine as _engine

    if db_engine is None:
        db_engine = _engine
    if session_factory is None:
        session_factory = _SessionLocal

    report = ImportReport()

    if not filepath.exists():
        print(f"ERROR: Archivo no encontrado: {filepath}")
        report.errors += 1
        report.errors_list.append({
            "message": f"Archivo no encontrado: {filepath}",
        })
        return report

    print(f"Leyendo archivo: {filepath}")
    sheets_data = read_excel(filepath)

    if not sheets_data:
        print("WARN: No se encontraron hojas con datos validos.")
        report.errors += 1
        report.errors_list.append({
            "message": "No se encontraron hojas con datos validos en el archivo.",
        })
        return report

    print(f"Hojas encontradas: {list(sheets_data.keys())}")

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    if "DESCUADRES" in wb.sheetnames:
        print("Procesando hoja DESCUADRES...")
        process_descadres(wb["DESCUADRES"], report, dry_run)
    wb.close()

    Base.metadata.create_all(bind=db_engine)

    all_records: list[dict] = []

    db = session_factory()
    try:
        for sheet_name, rows in sheets_data.items():
            if sheet_name.upper() == "DESCUADRES":
                continue

            period_date = parse_sheet_date(sheet_name)
            if period_date is None:
                report.sheets_skipped += 1
                report.skipped_sheets.append(sheet_name)
                report.warnings.append({
                    "sheet": sheet_name,
                    "row": None,
                    "product": "",
                    "message": f"No se pudo interpretar la fecha de la hoja '{sheet_name}'",
                })
                report.total_warnings += 1
                print(f"  Omitiendo hoja '{sheet_name}': no se pudo interpretar fecha")
                continue

            print(f"  Procesando hoja '{sheet_name}' (fecha: {period_date.strftime('%Y-%m-%d')})")
            records = process_sheet(db, sheet_name, rows, report, dry_run)
            all_records.extend(records)
            report.sheets_processed += 1

        for rec in all_records:
            if not dry_run:
                history = InventoryHistory(**rec)
                db.add(history)
            report.histories_created += 1

        if not dry_run:
            update_product_stock_and_ideal(db, all_records, report)

        if not dry_run:
            db.commit()
            print("\nCambios comprometidos (commit).")
        else:
            print("\nModo DRY RUN - no se modifico la base de datos.")

    except Exception as e:
        db.rollback()
        report.errors += 1
        report.errors_list.append({
            "message": f"Error critico: {e}",
        })
        print(f"\nERROR CRITICO: {e}")
        print("Rollback ejecutado.")
    finally:
        db.close()

    return report


def print_report(report: ImportReport, dry_run: bool) -> None:
    mode = "DRY RUN" if dry_run else "COMMIT"
    separator = "=" * 60

    print(f"\n{separator}")
    print("LA PATRONA - IMPORTACION HISTORICA")
    print(separator)
    print(f"Hojas procesadas:        {report.sheets_processed}")
    print(f"Hojas omitidas:          {report.sheets_skipped}")
    print(f"Filas leidas:            {report.rows_read}")
    print(f"Filas validas:           {report.rows_valid}")
    print(f"Filas omitidas:          {report.rows_skipped}")
    print(f"Productos encontrados:   {report.products_found}")
    print(f"Productos creados:       {report.products_created}")
    print(f"Productos pendientes:    {report.products_pending}")
    print(f"Historicos creados:      {report.histories_created}")
    print(f"Historicos existentes:   {report.histories_existing}")
    print(f"Advertencias inventario: {report.inventory_warnings}")
    print(f"Advertencias totales:    {report.total_warnings}")
    print(f"Errores:                 {report.errors}")
    print(f"Stock actual actualizado:{report.stock_updated} productos")
    print(f"Stock ideal actualizado: {report.ideal_stock_updated} productos")
    if report.descadres_rows > 0:
        print(f"DESCUADRES filas:        {report.descadres_rows}")
        print(f"DESCUADRES importadas:   {report.descadres_imported}")
    print(f"\nModo: {mode}")
    print(separator)

    if report.warnings:
        print(f"\nAdvertencias ({len(report.warnings)}):")
        for w in report.warnings[:20]:
            print(f"  [{w['sheet']}] Fila {w['row']}: {w['product']} - {w['message']}")
        if len(report.warnings) > 20:
            print(f"  ... y {len(report.warnings) - 20} mas")

    if report.errors_list:
        print(f"\nErrores ({len(report.errors_list)}):")
        for e in report.errors_list:
            print(f"  {e['message']}")

    if report.skipped_products:
        print(f"\nProductos pendientes ({len(report.skipped_products)}):")
        for p in report.skipped_products[:10]:
            print(f"  - {p}")
        if len(report.skipped_products) > 10:
            print(f"  ... y {len(report.skipped_products) - 10} mas")


def save_report(report: ImportReport, filepath: Path) -> None:
    report_path = filepath.parent / "historical_import_report.json"
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(report.to_dict(), f, indent=2, ensure_ascii=False, default=str)
    print(f"\nReporte guardado en: {report_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Importa historico de inventario y ventas desde Excel para La Patrona VIP."
    )
    parser.add_argument(
        "filepath",
        type=str,
        help="Ruta al archivo Excel (.xlsx) a importar.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Solo leer y validar sin modificar la base de datos.",
    )
    args = parser.parse_args()

    filepath = Path(args.filepath)
    report = run_import(filepath, dry_run=args.dry_run)
    print_report(report, dry_run=args.dry_run)
    save_report(report, filepath)


if __name__ == "__main__":
    main()
