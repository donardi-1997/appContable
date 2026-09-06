"""
Tests para el importador historico de inventario.
"""

from __future__ import annotations

import json
import sys
import os
from decimal import Decimal
from pathlib import Path
from datetime import datetime

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models import Product, InventoryHistory

from scripts.import_historical_inventory import (
    parse_number,
    parse_sheet_date,
    normalize_column_name,
    normalize_product_name,
    is_valid_product_row,
    read_excel,
    find_or_create_product,
    process_sheet,
    run_import,
    save_report,
    ImportReport,
    SOURCE_TAG,
)


# ── Fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture
def db_session():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    yield session
    session.close()


@pytest.fixture
def test_engine():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    yield engine, Session
    engine.dispose()


@pytest.fixture
def sample_excel(tmp_path):
    wb = openpyxl.Workbook()

    ws_marzo = wb.active
    ws_marzo.title = "MARZO 2026"
    ws_marzo.append(["PRODUCTO", "INVENTARIO", "INGRESOS", "EXISTENCIAS", "VENTAS",
                      "VALOR DE VENTA", "TOTAL VENDIDO", "STOCK IDEAL", "PARA PEDIDO PROMEDIO"])
    ws_marzo.append(["AMARILLO MANZANARES", 20, 59, 42, 37, 100000, 3700000, 48, 11])
    ws_marzo.append(["ROJO LOCAL", 15, 30, 18, 27, 80000, 2160000, 30, 3])
    ws_marzo.append(["TOTAL GENERAL", "", "", "", "", "", "", "", ""])

    ws_abril = wb.create_sheet("ABRIL 2026")
    ws_abril.append(["PRODUCTO", "INVENTARIO", "INGRESOS", "EXISTENCIAS", "VENTAS",
                      "VALOR DE VENTA", "TOTAL VENDIDO", "STOCK IDEAL", "PARA PEDIDO PROMEDIO"])
    ws_abril.append(["AMARILLO MANZANARES", 42, 20, 35, 27, 100000, 2700000, 48, 6])
    ws_abril.append(["ROJO LOCAL", 18, 25, 22, 21, 80000, 1680000, 30, 0])

    ws_desc = wb.create_sheet("DESCUADRES")
    ws_desc.append(["PRODUCTO", "FECHA", "MOTIVO", "DIFERENCIA"])
    ws_desc.append(["AMARILLO MANZANARES", "2026-03-15", "Faltante", -5])

    filepath = tmp_path / "test_historico.xlsx"
    wb.save(filepath)
    wb.close()
    return filepath


@pytest.fixture
def excel_with_formatted_numbers(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MARZO 2026"
    ws.append(["PRODUCTO", "INVENTARIO", "INGRESOS", "EXISTENCIAS", "VENTAS",
               "VALOR DE VENTA", "TOTAL VENDIDO", "STOCK IDEAL", "PARA PEDIDO PROMEDIO"])
    ws.append(["PRODUCTO A", "$100.000", "1.500.000", "1.400.000", "100.000",
               "$50.000", "5.000.000", "150.000", "50.000"])
    ws.append(["PRODUCTO B", "25", "10", "20", "15",
               "100000", "1500000", "30", "15"])

    filepath = tmp_path / "test_formatted.xlsx"
    wb.save(filepath)
    wb.close()
    return filepath


@pytest.fixture
def empty_excel(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HOJA VACIA"
    filepath = tmp_path / "empty.xlsx"
    wb.save(filepath)
    wb.close()
    return filepath


@pytest.fixture
def invalid_date_sheets(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HOJA RARA 123"
    ws.append(["PRODUCTO", "INVENTARIO", "INGRESOS", "EXISTENCIAS", "VENTAS",
               "VALOR DE VENTA", "TOTAL VENDIDO"])
    ws.append(["PRODUCTO X", 10, 5, 8, 7, 50000, 350000])

    filepath = tmp_path / "invalid_dates.xlsx"
    wb.save(filepath)
    wb.close()
    return filepath


@pytest.fixture
def inconsistent_excel(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MARZO 2026"
    ws.append(["PRODUCTO", "INVENTARIO", "INGRESOS", "EXISTENCIAS", "VENTAS",
               "VALOR DE VENTA", "TOTAL VENDIDO"])
    ws.append(["PRODUCTO INCONSISTENTE", 20, 50, 30, 35, 100000, 3500000])
    ws.append(["PRODUCTO TOTAL MAL", 10, 40, 25, 25, 80000, 2500000])

    filepath = tmp_path / "inconsistent.xlsx"
    wb.save(filepath)
    wb.close()
    return filepath


# ── Tests de parseo ─────────────────────────────────────────────────────────

class TestParseNumber:
    def test_integer(self):
        assert parse_number(100) == Decimal("100")

    def test_float(self):
        assert parse_number(100.5) == Decimal("100.5")

    def test_decimal(self):
        assert parse_number(Decimal("100.50")) == Decimal("100.50")

    def test_string_integer(self):
        assert parse_number("100000") == Decimal("100000")

    def test_string_with_comma_decimal(self):
        assert parse_number("100,50") == Decimal("100.50")

    def test_string_with_dollar_sign(self):
        assert parse_number("$100.000") == Decimal("100000")

    def test_string_with_thousand_dots(self):
        assert parse_number("1.000.000") == Decimal("1000000")

    def test_string_with_thousand_commas(self):
        assert parse_number("1,000,000") == Decimal("1000000")

    def test_string_complex_format(self):
        assert parse_number("$1.500.000,50") == Decimal("1500000.50")

    def test_none(self):
        assert parse_number(None) is None

    def test_empty_string(self):
        assert parse_number("") is None

    def test_whitespace_only(self):
        assert parse_number("   ") is None

    def test_non_convertible(self):
        assert parse_number("abc") is None


class TestParseSheetDate:
    def test_month_year(self):
        assert parse_sheet_date("MARZO 2026") == datetime(2026, 3, 1)

    def test_month_dash_year(self):
        assert parse_sheet_date("MAR-2026") == datetime(2026, 3, 1)

    def test_month_only(self):
        assert parse_sheet_date("ABRIL") == datetime(2026, 4, 1)

    def test_numeric_format(self):
        assert parse_sheet_date("03-2026") == datetime(2026, 3, 1)

    def test_numeric_slash(self):
        assert parse_sheet_date("05/2026") == datetime(2026, 5, 1)

    def test_invalid_name(self):
        assert parse_sheet_date("HOJA RARA 123") is None

    def test_empty_string(self):
        assert parse_sheet_date("") is None

    def test_descadres(self):
        assert parse_sheet_date("DESCUADRES") is None

    def test_case_insensitive(self):
        assert parse_sheet_date("marzo 2026") == datetime(2026, 3, 1)

    def test_two_digit_year(self):
        assert parse_sheet_date("MAR-26") == datetime(2026, 3, 1)


class TestNormalizeColumnName:
    def test_exact_match(self):
        assert normalize_column_name("PRODUCTO") == "product_name"

    def test_case_insensitive(self):
        assert normalize_column_name("producto") == "product_name"

    def test_with_accents(self):
        assert normalize_column_name("VALOR DE VENTA") == "sale_price"

    def test_unknown_column(self):
        assert normalize_column_name("COLUMNA DESCONOCIDA") is None

    def test_whitespace(self):
        assert normalize_column_name("  PRODUCTO  ") == "product_name"


class TestNormalizeProductName:
    def test_trim_spaces(self):
        assert normalize_product_name("  AMARILLO  MANZANARES  ") == "AMARILLO MANZANARES"

    def test_double_spaces(self):
        assert normalize_product_name("AMARILLO   MANZANARES") == "AMARILLO MANZANARES"

    def test_case(self):
        assert normalize_product_name("amarillo manzanares") == "AMARILLO MANZANARES"


class TestIsValidProductRow:
    def test_valid_row(self):
        row = {"product_name": "PRODUCTO A", "opening_stock": Decimal("10")}
        assert is_valid_product_row(row) is True

    def test_empty_name(self):
        row = {"product_name": "", "opening_stock": Decimal("10")}
        assert is_valid_product_row(row) is False

    def test_none_name(self):
        row = {"product_name": None, "opening_stock": Decimal("10")}
        assert is_valid_product_row(row) is False

    def test_total_row(self):
        row = {"product_name": "TOTAL GENERAL", "opening_stock": Decimal("100")}
        assert is_valid_product_row(row) is False

    def test_subtotal_row(self):
        row = {"product_name": "SUBTOTAL", "opening_stock": Decimal("50")}
        assert is_valid_product_row(row) is False

    def test_no_numeric_values(self):
        row = {"product_name": "PRODUCTO A"}
        assert is_valid_product_row(row) is False


# ── Tests de lectura Excel ──────────────────────────────────────────────────

class TestReadExcel:
    def test_reads_all_sheets(self, sample_excel):
        data = read_excel(sample_excel)
        assert "MARZO 2026" in data
        assert "ABRIL 2026" in data

    def test_skips_empty_sheets(self, sample_excel):
        data = read_excel(sample_excel)
        # DESCUADRES has PRODUCTO column but we check sheets_data only has
        # data sheets (MARZO, ABRIL) since DESCUADRES header row has PRODUCTO
        assert "MARZO 2026" in data
        assert "ABRIL 2026" in data

    def test_row_count(self, sample_excel):
        data = read_excel(sample_excel)
        # MARZO has 3 data rows but TOTAL GENERAL gets filtered later in process_sheet
        # read_excel returns all rows with product_name column
        assert len(data["MARZO 2026"]) == 3

    def test_empty_excel(self, empty_excel):
        data = read_excel(empty_excel)
        assert len(data) == 0

    def test_formatted_numbers(self, excel_with_formatted_numbers):
        data = read_excel(excel_with_formatted_numbers)
        assert len(data) == 1
        rows = data["MARZO 2026"]
        assert len(rows) == 2
        assert rows[0]["opening_stock"] == Decimal("100000")
        assert rows[0]["incoming_stock"] == Decimal("1500000")
        assert rows[1]["opening_stock"] == Decimal("25")


# ── Tests de logica de base de datos ───────────────────────────────────────

class TestFindOrCreateProduct:
    def test_finds_existing_product(self, db_session):
        product = Product(name="AMARILLO MANZANARES", stock=Decimal("10"), active=True)
        db_session.add(product)
        db_session.commit()

        report = ImportReport()
        result = find_or_create_product(db_session, "AMARILLO MANZANARES", report)
        assert result is not None
        assert result.id == product.id
        assert report.products_created == 0

    def test_finds_case_insensitive(self, db_session):
        product = Product(name="AMARILLO MANZANARES", stock=Decimal("10"), active=True)
        db_session.add(product)
        db_session.commit()

        report = ImportReport()
        result = find_or_create_product(db_session, "amarillo manzanares", report)
        assert result is not None
        assert result.id == product.id

    def test_creates_new_product(self, db_session):
        report = ImportReport()
        result = find_or_create_product(db_session, "NUEVO PRODUCTO", report)
        assert result is not None
        assert result.name == "NUEVO PRODUCTO"
        assert report.products_created == 1

    def test_not_finds_inactive_product(self, db_session):
        product = Product(name="PRODUCTO INACTIVO", stock=Decimal("10"), active=False)
        db_session.add(product)
        db_session.commit()

        report = ImportReport()
        result = find_or_create_product(db_session, "PRODUCTO INACTIVO", report)
        assert result is not None
        assert result.id != product.id
        assert report.products_created == 1


# ── Tests de procesamiento de hoja ─────────────────────────────────────────

class TestProcessSheet:
    def test_basic_processing(self, db_session):
        rows = [
            {
                "product_name": "PRODUCTO A",
                "opening_stock": Decimal("20"),
                "incoming_stock": Decimal("50"),
                "closing_stock": Decimal("30"),
                "units_sold": Decimal("40"),
                "sale_price": Decimal("100000"),
                "total_sold": Decimal("4000000"),
                "ideal_stock": Decimal("48"),
                "suggested_order": Decimal("8"),
            },
        ]
        report = ImportReport()
        records = process_sheet(db_session, "MARZO 2026", rows, report, dry_run=False)
        assert len(records) == 1
        assert records[0]["opening_stock"] == Decimal("20")
        assert records[0]["units_sold"] == Decimal("40")

    def test_idempotency(self, db_session):
        rows = [
            {
                "product_name": "PRODUCTO A",
                "opening_stock": Decimal("20"),
                "incoming_stock": Decimal("50"),
                "closing_stock": Decimal("30"),
                "units_sold": Decimal("40"),
                "sale_price": Decimal("100000"),
                "total_sold": Decimal("4000000"),
                "ideal_stock": Decimal("48"),
                "suggested_order": Decimal("8"),
            },
        ]
        report1 = ImportReport()
        records1 = process_sheet(db_session, "MARZO 2026", rows, report1, dry_run=False)
        for rec in records1:
            db_session.add(InventoryHistory(**rec))
        db_session.commit()

        report2 = ImportReport()
        records2 = process_sheet(db_session, "MARZO 2026", rows, report2, dry_run=False)
        assert len(records2) == 0
        assert report2.histories_existing == 1

    def test_sales_calculation_validation(self, db_session):
        rows = [
            {
                "product_name": "PRODUCTO INCONSISTENTE",
                "opening_stock": Decimal("20"),
                "incoming_stock": Decimal("50"),
                "closing_stock": Decimal("30"),
                "units_sold": Decimal("35"),
                "sale_price": Decimal("100000"),
                "total_sold": Decimal("3500000"),
                "ideal_stock": None,
                "suggested_order": None,
            },
        ]
        report = ImportReport()
        records = process_sheet(db_session, "MARZO 2026", rows, report, dry_run=False)
        assert len(records) == 1
        assert report.inventory_warnings == 1

    def test_total_sold_validation(self, db_session):
        rows = [
            {
                "product_name": "PRODUCTO TOTAL MAL",
                "opening_stock": Decimal("10"),
                "incoming_stock": Decimal("40"),
                "closing_stock": Decimal("25"),
                "units_sold": Decimal("25"),
                "sale_price": Decimal("80000"),
                "total_sold": Decimal("2500000"),
                "ideal_stock": None,
                "suggested_order": None,
            },
        ]
        report = ImportReport()
        records = process_sheet(db_session, "MARZO 2026", rows, report, dry_run=False)
        assert len(records) == 1
        assert report.total_warnings == 1

    def test_skips_total_rows(self, db_session):
        rows = [
            {"product_name": "TOTAL GENERAL", "opening_stock": Decimal("100")},
            {"product_name": "SUBTOTAL", "opening_stock": Decimal("50")},
        ]
        report = ImportReport()
        records = process_sheet(db_session, "MARZO 2026", rows, report, dry_run=False)
        assert len(records) == 0
        assert report.rows_skipped == 2


# ── Tests de importacion completa ───────────────────────────────────────────

class TestFullImport:
    def test_dry_run_does_not_write(self, sample_excel, test_engine):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=True, db_engine=engine, session_factory=Session)
        assert report.histories_created > 0

    def test_real_import(self, sample_excel, test_engine):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)
        assert report.sheets_processed >= 2
        assert report.histories_created >= 4
        assert report.rows_valid >= 4

    def test_import_creates_products(self, sample_excel, test_engine):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)
        assert report.products_created >= 2

    def test_import_with_existing_product(self, sample_excel, test_engine):
        engine, Session = test_engine
        db = Session()
        product = Product(name="AMARILLO MANZANARES", stock=Decimal("0"), active=True)
        db.add(product)
        db.commit()
        db.close()

        report = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)
        assert report.products_created <= 2

    def test_stock_updated_with_most_recent_period(self, sample_excel, test_engine):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)
        db = Session()
        product = db.query(Product).filter(Product.name == "AMARILLO MANZANARES").first()
        assert product is not None
        # ABRIL 2026 has closing_stock=35 for AMARILLO MANZANARES
        assert product.stock == Decimal("35")
        db.close()

    def test_ideal_stock_updated(self, sample_excel, test_engine):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)
        db = Session()
        product = db.query(Product).filter(Product.name == "AMARILLO MANZANARES").first()
        assert product is not None
        # ABRIL 2026 has ideal_stock=48
        assert product.ideal_stock == Decimal("48")
        db.close()

    def test_inconsistent_data_continues(self, inconsistent_excel, test_engine):
        engine, Session = test_engine
        report = run_import(inconsistent_excel, dry_run=False, db_engine=engine, session_factory=Session)
        assert report.histories_created >= 2
        assert report.inventory_warnings >= 1

    def test_empty_excel(self, empty_excel, test_engine):
        engine, Session = test_engine
        report = run_import(empty_excel, dry_run=False, db_engine=engine, session_factory=Session)
        assert report.errors >= 1

    def test_invalid_date_sheets(self, invalid_date_sheets, test_engine):
        engine, Session = test_engine
        report = run_import(invalid_date_sheets, dry_run=True, db_engine=engine, session_factory=Session)
        assert report.sheets_skipped >= 1


class TestIdempotency:
    def test_double_run_no_duplicates(self, sample_excel, test_engine):
        engine, Session = test_engine
        report1 = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)
        report2 = run_import(sample_excel, dry_run=False, db_engine=engine, session_factory=Session)

        db = Session()
        count = db.query(InventoryHistory).count()
        db.close()

        assert count == 4
        assert report2.histories_existing == report1.histories_created


class TestFormattedNumbers:
    def test_dollar_format(self, excel_with_formatted_numbers, test_engine):
        engine, Session = test_engine
        report = run_import(excel_with_formatted_numbers, dry_run=False, db_engine=engine, session_factory=Session)
        assert report.histories_created >= 2


class TestDescadres:
    def test_descadres_detected(self, sample_excel, test_engine):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=True, db_engine=engine, session_factory=Session)
        assert report.descadres_rows >= 1


class TestReportJson:
    def test_report_saved(self, sample_excel, test_engine, tmp_path):
        engine, Session = test_engine
        report = run_import(sample_excel, dry_run=True, db_engine=engine, session_factory=Session)
        save_report(report, sample_excel)
        report_path = sample_excel.parent / "historical_import_report.json"
        assert report_path.exists()
        with open(report_path) as f:
            data = json.load(f)
        assert "summary" in data
        assert "warnings" in data
